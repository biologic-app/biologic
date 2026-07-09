"""Direction + samples import (Excel primary, JSON fallback).

Both importers accept the same flat "one row per sample" shape: every row
carries direction-level fields (repeated across the rows that belong to the
same direction) plus sample-level fields. Rows are grouped into directions by
the ``(year_no, base_no)`` pair — the first row seen for a given pair supplies
the direction fields, every row for that pair adds one sample. Imported
directions and samples land in their default lifecycle status (``draft`` /
``pending``); the registrar edits and registers them afterwards.
"""

from __future__ import annotations

import json
from collections.abc import Callable
from dataclasses import dataclass, field
from datetime import datetime
from io import BytesIO
from typing import Any
from uuid import UUID

from openpyxl import load_workbook
from pydantic import BaseModel

from src.core.errors import ValidationError

MAX_IMPORT_BYTES = 5 * 1024 * 1024

_DIRECTION_FIELDS = (
    "year_no",
    "base_no",
    "is_urgent",
    "doctor_id",
    "object_id",
    "sampled_at",
    "received_at",
    "completed_at",
)

_SAMPLE_FIELDS = (
    "sample_name",
    "sample_type_id",
    "mass",
    "alternate_name",
    "nomenclature_code",
    "batch_code",
    "supplier",
    "sample_is_urgent",
    "month_no",
)


def _parse_int(value: Any) -> int:
    return int(value)


def _parse_bool(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    normalized = str(value).strip().lower()
    if normalized in {"1", "true", "yes", "y", "да"}:
        return True
    if normalized in {"0", "false", "no", "n", "нет"}:
        return False
    raise ValueError("Invalid boolean value.")


def _parse_uuid(value: Any) -> UUID:
    return UUID(str(value))


def _parse_datetime(value: Any) -> datetime:
    if isinstance(value, datetime):
        return value
    return datetime.fromisoformat(str(value).replace("Z", "+00:00"))


def _parse_str(value: Any) -> str:
    return str(value)


_FIELD_PARSERS: dict[str, Callable[[Any], object]] = {
    "year_no": _parse_int,
    "base_no": _parse_int,
    "is_urgent": _parse_bool,
    "doctor_id": _parse_uuid,
    "object_id": _parse_uuid,
    "sampled_at": _parse_datetime,
    "received_at": _parse_datetime,
    "completed_at": _parse_datetime,
    "sample_name": _parse_str,
    "sample_type_id": _parse_uuid,
    "mass": _parse_str,
    "alternate_name": _parse_str,
    "nomenclature_code": _parse_str,
    "batch_code": _parse_str,
    "supplier": _parse_str,
    "sample_is_urgent": _parse_bool,
    "month_no": _parse_int,
}

_ALL_FIELDS = _DIRECTION_FIELDS + _SAMPLE_FIELDS


class WorkflowImportSummary(BaseModel):
    filename: str
    rows_processed: int
    directions_created: int
    samples_created: int
    skipped_rows: int
    errors: list[dict[str, object]]
    warnings: list[dict[str, object]]


def _issue(row: int | None, field: str, message: str) -> dict[str, object]:
    payload: dict[str, object] = {"field": field, "message": message}
    if row is not None:
        payload["row"] = row
    return payload


def _parse_row(row: dict[str, Any]) -> tuple[dict[str, object], list[dict[str, object]]]:
    values: dict[str, object] = {}
    errors: list[dict[str, object]] = []

    for field_name in _ALL_FIELDS:
        raw = row.get(field_name)
        if raw is None or (isinstance(raw, str) and raw.strip() == ""):
            if field_name in ("year_no", "sample_name"):
                errors.append(_issue(None, field_name, "Value is required."))
            continue
        try:
            values[field_name] = _FIELD_PARSERS[field_name](
                raw.strip() if isinstance(raw, str) else raw
            )
        except (TypeError, ValueError):
            errors.append(_issue(None, field_name, "Value has invalid format."))

    return values, errors


@dataclass
class _PendingDirection:
    fields: dict[str, object]
    samples: list[dict[str, object]] = field(default_factory=list)
    row_numbers: list[int] = field(default_factory=list)


_SAMPLE_COPY_FIELDS = (
    ("sample_type_id", "sample_type_id"),
    ("mass", "mass"),
    ("alternate_name", "alternate_name"),
    ("nomenclature_code", "nomenclature_code"),
    ("batch_code", "batch_code"),
    ("supplier", "supplier"),
    ("month_no", "month_no"),
)


class DirectionSampleImporter:
    """Groups parsed rows into directions with nested samples and persists them."""

    def __init__(self, *, directions: Any, samples: Any, created_by: UUID | None = None) -> None:
        self.directions = directions
        self.samples = samples
        self.created_by = created_by

    async def import_rows(
        self, filename: str, rows: list[dict[str, Any]], warnings: list[dict[str, object]]
    ) -> WorkflowImportSummary:
        errors: list[dict[str, object]] = []
        directions_by_key: dict[tuple[int, int | None], _PendingDirection] = {}
        order: list[tuple[int, int | None]] = []
        skipped_rows = 0

        for row_number, raw_row in enumerate(rows, start=2):
            values, row_errors = _parse_row(raw_row)
            if row_errors:
                errors.append({"row": row_number, "errors": row_errors})
                skipped_rows += 1
                continue
            if "sample_name" not in values:
                sample_name_issue = _issue(None, "sample_name", "Value is required.")
                errors.append({"row": row_number, "errors": [sample_name_issue]})
                skipped_rows += 1
                continue

            year_no = values["year_no"]
            base_no = values.get("base_no")
            assert isinstance(year_no, int)
            assert base_no is None or isinstance(base_no, int)
            key = (year_no, base_no)
            if key not in directions_by_key:
                direction_fields = {
                    dfield: values[dfield] for dfield in _DIRECTION_FIELDS if dfield in values
                }
                directions_by_key[key] = _PendingDirection(fields=direction_fields)
                order.append(key)

            pending = directions_by_key[key]
            pending.row_numbers.append(row_number)
            sample_values: dict[str, object] = {"name": values["sample_name"]}
            for source_field, target_field in _SAMPLE_COPY_FIELDS:
                if source_field in values:
                    sample_values[target_field] = values[source_field]
            sample_values["is_urgent"] = values.get(
                "sample_is_urgent", pending.fields.get("is_urgent", False)
            )
            pending.samples.append(sample_values)

        directions_created = 0
        samples_created = 0
        for key in order:
            pending = directions_by_key[key]
            year_no, base_no = key
            existing = await self.directions.find_by_year_and_base_no(year_no, base_no)
            if existing is not None:
                skipped_rows += len(pending.row_numbers)
                duplicate_issue = _issue(
                    None,
                    "year_no",
                    "Направление с таким годом и номером уже существует, пропущено.",
                )
                for row_number in pending.row_numbers:
                    errors.append({"row": row_number, "errors": [duplicate_issue]})
                continue

            direction_row = await self.directions.create(
                pending.fields, created_by=self.created_by
            )
            directions_created += 1
            for sample_values in pending.samples:
                await self.samples.create({**sample_values, "direction_id": direction_row.id})
                samples_created += 1

        return WorkflowImportSummary(
            filename=filename,
            rows_processed=len(rows),
            directions_created=directions_created,
            samples_created=samples_created,
            skipped_rows=skipped_rows,
            errors=errors,
            warnings=warnings,
        )


class DirectionExcelImportService:
    """Imports directions + samples from an .xlsx workbook (first worksheet)."""

    def __init__(self, *, directions: Any, samples: Any, created_by: UUID | None = None) -> None:
        self._importer = DirectionSampleImporter(
            directions=directions, samples=samples, created_by=created_by
        )

    async def import_file(self, filename: str, content: bytes) -> WorkflowImportSummary:
        if not filename.lower().endswith(".xlsx"):
            raise ValidationError(
                "Only .xlsx direction imports are supported.",
                extra={"errors": [_issue(None, "file", "Only .xlsx files are supported.")]},
            )
        if len(content) > MAX_IMPORT_BYTES:
            raise ValidationError(
                "Direction import file is too large.",
                extra={"errors": [_issue(None, "file", "File must not exceed 5 MB.")]},
            )

        try:
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001 — any parse failure is a 422
            raise ValidationError(
                "Direction import file is not a valid .xlsx workbook.",
                extra={"errors": [_issue(None, "file", "Workbook could not be parsed.")]},
            ) from exc

        sheet = workbook.worksheets[0]
        rows_iter = sheet.iter_rows(values_only=True)
        try:
            header = next(rows_iter)
        except StopIteration:
            raise ValidationError(
                "Direction import workbook has no header row.",
                extra={"errors": [_issue(1, "header", "Header row is required.")]},
            ) from None

        fieldnames = [str(cell).strip() if cell is not None else "" for cell in header]
        if "year_no" not in fieldnames:
            raise ValidationError(
                "Direction import workbook must contain year_no.",
                extra={"errors": [_issue(1, "year_no", "year_no column is required.")]},
            )

        unknown_fields = sorted(set(fieldnames) - set(_ALL_FIELDS) - {""})
        warnings = [_issue(None, field, "Column is ignored.") for field in unknown_fields]

        rows: list[dict[str, Any]] = []
        for raw_row in rows_iter:
            if raw_row is None or all(cell is None for cell in raw_row):
                continue
            rows.append(dict(zip(fieldnames, raw_row, strict=False)))

        return await self._importer.import_rows(filename, rows, warnings)


class DirectionJsonImportService:
    """Fallback importer: accepts the same row shape as Excel, in JSON form.

    Body shape: ``{"rows": [{...row...}, ...]}`` or a bare JSON array of rows.
    """

    def __init__(self, *, directions: Any, samples: Any, created_by: UUID | None = None) -> None:
        self._importer = DirectionSampleImporter(
            directions=directions, samples=samples, created_by=created_by
        )

    async def import_file(self, filename: str, content: bytes) -> WorkflowImportSummary:
        if not filename.lower().endswith(".json"):
            raise ValidationError(
                "Only .json direction imports are supported.",
                extra={"errors": [_issue(None, "file", "Only .json files are supported.")]},
            )
        if len(content) > MAX_IMPORT_BYTES:
            raise ValidationError(
                "Direction import file is too large.",
                extra={"errors": [_issue(None, "file", "File must not exceed 5 MB.")]},
            )

        try:
            payload = json.loads(content.decode("utf-8-sig"))
        except (UnicodeDecodeError, json.JSONDecodeError) as exc:
            raise ValidationError(
                "Direction import file must be valid UTF-8 JSON.",
                extra={"errors": [_issue(None, "file", "File must be valid JSON.")]},
            ) from exc

        rows = payload.get("rows") if isinstance(payload, dict) else payload
        if not isinstance(rows, list):
            raise ValidationError(
                "Direction import JSON must contain a list of rows.",
                extra={"errors": [_issue(None, "rows", "A JSON array of rows is required.")]},
            )

        fieldnames = {key for row in rows if isinstance(row, dict) for key in row}
        unknown_fields = sorted(fieldnames - set(_ALL_FIELDS))
        warnings = [_issue(None, field, "Field is ignored.") for field in unknown_fields]

        return await self._importer.import_rows(filename, rows, warnings)
