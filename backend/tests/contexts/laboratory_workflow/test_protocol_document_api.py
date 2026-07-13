from collections.abc import Iterator
from datetime import UTC, datetime
from io import BytesIO
from uuid import UUID

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pytest import MonkeyPatch

from src.app_factory import create_app
from src.contexts.laboratory_workflow.application.protocol_document import (
    ProtocolDocumentData,
    ProtocolSampleRow,
    protocol_document_filename,
    protocol_excerpt_filename,
    render_excerpt,
    render_full_document,
)
from src.contexts.laboratory_workflow.presentation.router import (
    get_workflow_crud_use_case,
)
from src.core.config import get_settings
from src.core.errors import NotFoundError

_XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_PROTOCOL_ID = UUID("0192f4a0-0000-7000-8000-000000000001")
_MISSING_ID = UUID("0192f4a0-0000-7000-8000-0000000000ff")


def _document_data() -> ProtocolDocumentData:
    return ProtocolDocumentData(
        protocol_id=_PROTOCOL_ID,
        year_no=2026,
        formed_at=datetime(2026, 7, 1, 9, 0, tzinfo=UTC),
        protocol_type="Испытательный",
        conclusion="Соответствует",
        copies=2,
        rows=(
            ProtocolSampleRow(
                name="Проба 1",
                sample_type="Вода",
                status="Брак",
                direction_no="2026-15",
                received_at=None,
                reject_reason="Повреждена тара",
            ),
        ),
    )


class FakeWorkflowCrudUseCase:
    async def protocol_document(self, protocol_id: UUID) -> tuple[str, bytes]:
        if protocol_id == _MISSING_ID:
            raise NotFoundError(f"protocols item {protocol_id} was not found.")
        data = _document_data()
        return protocol_document_filename(data), render_full_document(data)

    async def protocol_excerpt(self, protocol_id: UUID) -> tuple[str, bytes]:
        if protocol_id == _MISSING_ID:
            raise NotFoundError(f"protocols item {protocol_id} was not found.")
        data = _document_data()
        return protocol_excerpt_filename(data), render_excerpt(data)


@pytest.fixture
def client(monkeypatch: MonkeyPatch) -> Iterator[TestClient]:
    monkeypatch.setenv("APP_DATABASE_URL", "postgresql+asyncpg://user:pass@localhost:5432/test")
    monkeypatch.setenv("APP_JWT_SECRET_KEY", "test-secret")
    monkeypatch.setenv("APP_AUTH_COOKIE_SECURE", "false")
    monkeypatch.setenv("APP_AUTH_COOKIE_DOMAIN", "localhost")
    get_settings.cache_clear()
    try:
        app = create_app()
        app.dependency_overrides[get_workflow_crud_use_case] = FakeWorkflowCrudUseCase
        yield TestClient(app)
    finally:
        get_settings.cache_clear()


def test_protocol_document_returns_xlsx(client: TestClient) -> None:
    response = client.get(f"/api/v1/protocols/{_PROTOCOL_ID}/document")

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(_XLSX_MEDIA_TYPE)
    assert response.headers["content-disposition"] == (
        'attachment; filename="protocol_2026-0192f4a0.xlsx"'
    )
    worksheet = load_workbook(BytesIO(response.content)).active
    rows = list(worksheet.iter_rows(values_only=True))
    assert rows[0][0] == "Протокол испытаний"


def test_protocol_excerpt_returns_xlsx(client: TestClient) -> None:
    response = client.get(f"/api/v1/protocols/{_PROTOCOL_ID}/excerpt")

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(_XLSX_MEDIA_TYPE)
    assert response.headers["content-disposition"] == (
        'attachment; filename="protocol_2026-0192f4a0_excerpt.xlsx"'
    )
    worksheet = load_workbook(BytesIO(response.content)).active
    rows = list(worksheet.iter_rows(values_only=True))
    assert rows[0][0] == "Выписка из протокола — бракованные образцы"
    assert any("Причина брака" in row for row in rows)


def test_protocol_document_missing_protocol_returns_404(client: TestClient) -> None:
    response = client.get(f"/api/v1/protocols/{_MISSING_ID}/document")

    assert response.status_code == 404
