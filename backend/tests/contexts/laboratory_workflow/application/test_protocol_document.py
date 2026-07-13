from datetime import UTC, datetime
from io import BytesIO
from uuid import UUID

from openpyxl import load_workbook

from src.contexts.laboratory_workflow.application.protocol_document import (
    ProtocolDocumentData,
    ProtocolSampleRow,
    protocol_document_filename,
    protocol_excerpt_filename,
    render_excerpt,
    render_full_document,
)

_PROTOCOL_ID = UUID("0192f4a0-0000-7000-8000-000000000001")


def _data(rows: tuple[ProtocolSampleRow, ...]) -> ProtocolDocumentData:
    return ProtocolDocumentData(
        protocol_id=_PROTOCOL_ID,
        year_no=2026,
        formed_at=datetime(2026, 7, 1, 9, 0, tzinfo=UTC),
        protocol_type="Испытательный",
        conclusion="Соответствует",
        copies=3,
        rows=rows,
    )


def _sheet_rows(content: bytes) -> list[tuple[object, ...]]:
    workbook = load_workbook(BytesIO(content))
    worksheet = workbook.active
    return list(worksheet.iter_rows(values_only=True))


def _table(content: bytes) -> tuple[tuple[object, ...], list[tuple[object, ...]]]:
    rows = _sheet_rows(content)
    header_index = next(
        index for index, row in enumerate(rows) if row and row[0] == "Название образца"
    )
    header = rows[header_index]
    data_rows = [row for row in rows[header_index + 1 :] if row and row[0] is not None]
    return header, data_rows


def test_render_full_document_produces_parseable_xlsx_with_rows() -> None:
    data = _data(
        (
            ProtocolSampleRow(
                name="Проба 1",
                sample_type="Вода",
                status="Завершён",
                direction_no="2026-15",
                received_at=datetime(2026, 6, 1, 12, 0, tzinfo=UTC),
            ),
            ProtocolSampleRow(
                name="Проба 2",
                sample_type="Почва",
                status="Брак",
                direction_no="2026-16",
                received_at=None,
            ),
        ),
    )

    content = render_full_document(data)

    header, data_rows = _table(content)
    assert header == (
        "Название образца",
        "Тип образца",
        "Статус",
        "Направление",
        "Дата поступления",
    )
    assert len(data_rows) == 2
    assert data_rows[0][0] == "Проба 1"
    assert _sheet_rows(content)[0][0] == "Протокол испытаний"


def test_render_excerpt_includes_reason_column() -> None:
    data = _data(
        (
            ProtocolSampleRow(
                name="Проба брак",
                sample_type="Вода",
                status="Брак",
                direction_no="2026-17",
                received_at=None,
                reject_reason="Повреждена тара",
            ),
        ),
    )

    content = render_excerpt(data)

    header, data_rows = _table(content)
    assert header[-1] == "Причина брака"
    assert len(data_rows) == 1
    assert data_rows[0][-1] == "Повреждена тара"
    assert _sheet_rows(content)[0][0] == "Выписка из протокола — бракованные образцы"


def test_render_document_with_no_rows_is_valid_and_empty() -> None:
    content = render_full_document(_data(()))

    header, data_rows = _table(content)
    assert header[0] == "Название образца"
    assert data_rows == []


def test_document_filenames_use_protocol_number() -> None:
    data = _data(())
    assert protocol_document_filename(data) == "protocol_2026-0192f4a0.xlsx"
    assert protocol_excerpt_filename(data) == "protocol_2026-0192f4a0_excerpt.xlsx"
